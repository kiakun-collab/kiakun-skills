"""Contract fixtures test orchestration, not image recognition or visual quality."""
from __future__ import annotations
import json
import sys
import tempfile
import unittest
from pathlib import Path
from PIL import Image
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import rebuild_workflow as workflow
from validate_rebuild_evidence import validate_geometry_exceptions
from audit_image_alpha import audit_assets
from tests.common import make_basic_pptx


class WorkflowTests(unittest.TestCase):
    def create_task(self, root, extra=None):
        task = {"executionProfile": "balanced", "routingIntent": {"operation": "rebuild", "deliverable": "editable"},
                "userCopy": [{"id": "body", "page": 1, "text": "测试正文", "source": "test user instruction",
                              "pptxShapeName": "body-text-main"}]}
        task.update(extra or {})
        workflow.write(root / "task.json", task)
        workflow.initialize(root / "task.json", root / "run")
        return task

    def fixture(self, root, extra=None):
        task = self.create_task(root, extra)
        deck = make_basic_pptx(root / "deck.pptx")
        image = root / "slide-1.png"
        Image.new("RGB", (320, 180), "white").save(image)
        workflow.write(root / "receipt.json", {"outputPptxSha256": workflow.sha(deck), "acceptanceRenderer": "test fixture only",
            "pages": [{"page": 1, "renderPath": str(image), "renderSha256": workflow.sha(image)}]})
        result = workflow.prepare_review(root / "run", deck, root / "receipt.json")
        review_path = Path(result["review"])
        review = workflow.read(review_path)
        review.update(contentStatus="PASS", visionAuditStatus="PASS", visualOverlapCount=0,
                      visualFidelityStatus="PASS", majorFidelityDeviationCount=0, minorFidelityDeviationCount=0,
                      visibleAssetSeamCount=0, editableBoundaryStatus="PASS", autoFidelityBlocked=False)
        review["pages"][0].update(reviewer="test fixture", observations="Simulated review to exercise field assembly",
                                  coordinateStatus="PASS", typographyStatus="PASS")
        for gate in review.get("level3Gates", {}).values():
            gate.update(status="PASS", automatedEvidence="fixture operation", manualEvidence="fixture review")
        workflow.write(review_path, review)
        workflow.write(root / "extraction.json", {"textBlocks": [{"id": "body", "text": "测试正文"}]})
        workflow.write(root / "layout.json", {"objects": [{"name": "body-text-main", "sourceExtractionId": "body"}]})
        workflow.write(root / "pairing.json", {"renderDirectory": str(root), "pairings": [{"page": 1, "render": image.name}]})
        evidence = {"reviewFile": str(review_path), "visualExtractionFiles": ["extraction.json"],
                    "layoutSpecFiles": ["layout.json"], "referenceRenderComparison": image.name,
                    "referenceRenderPairingManifest": "pairing.json"}
        workflow.write(root / "evidence.json", evidence)
        event = workflow.begin(root / "run", "build", [1])
        workflow.end(root / "run", event["id"], 0)
        return task, review_path

    def test_user_copy_overrides_only_matching_text_and_preserves_geometry(self):
        task = {"userCopy": [{"id": "title", "page": 1, "text": "新标题 120%", "source": "user"}]}
        extraction = {"page": 1, "textBlocks": [{"id": "title", "text": "图中旧字", "bbox": {"x": 10}},
                                                {"id": "body", "text": "未提供的正文"}]}
        result = workflow.bind_copy(task, extraction)
        self.assertEqual(result["textBlocks"][0]["text"], "新标题 120%")
        self.assertFalse(result["textBlocks"][0]["requiresTextRecognition"])
        self.assertEqual(result["textBlocks"][0]["bbox"], {"x": 10})
        self.assertEqual(result["textBlocks"][1]["text"], "未提供的正文")
        self.assertEqual(extraction["textBlocks"][0]["text"], "图中旧字")

    def test_route_snapshot_is_required_and_immutable(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            self.create_task(root)
            task = workflow.read(root / "run/task-input.json")
            task["executionProfile"] = "fast"
            workflow.write(root / "run/task-input.json", task)
            with self.assertRaisesRegex(ValueError, "changed after routing"):
                workflow.begin(root / "run", "build", [1])

    def test_completed_run_directory_is_not_overwritten(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            self.create_task(root)
            with self.assertRaises(ValueError):
                workflow.initialize(root / "task.json", root / "run")

    def test_command_failures_keep_distinct_logs_and_environment(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            self.create_task(root, {"runtime": {"nodeModules": "fixture-modules"}})
            failed = workflow.run_command(root / "run", "build", [1], 0, 1,
                [sys.executable, "-c", "raise SystemExit(3)"])
            passed = workflow.run_command(root / "run", "build", [1], 0, 1,
                [sys.executable, "-c", "import os; assert os.environ['RUNTIME_NODE_MODULES']=='fixture-modules'"])
            self.assertEqual(failed["exitCode"], 3)
            self.assertEqual(passed["exitCode"], 0)
            self.assertNotEqual(failed["log"], passed["log"])
            self.assertEqual(workflow.timings(root / "run")["autoIterationCount"], 0)

    def test_budget_and_event_identity(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            self.create_task(root)
            with self.assertRaises(ValueError): workflow.begin(root / "run", "build", [1], 3)
            event = workflow.begin(root / "run", "imagegen", [1])
            workflow.end(root / "run", event["id"], 1)
            with self.assertRaises(ValueError): workflow.end(root / "run", event["id"], 0)
            self.assertEqual(workflow.timings(root / "run")["imageGenerationCount"], 1)

    def test_design_wait_or_reference_only_cannot_start_build(self):
        for intent in ({"redesignRequested": True, "userSelectsDesign": True},
                       {"redesignRequested": True, "referenceOnly": True}):
            with tempfile.TemporaryDirectory() as tmp:
                root = Path(tmp)
                self.create_task(root, {"routingIntent": intent})
                with self.assertRaisesRegex(ValueError, "not authorized"):
                    workflow.begin(root / "run", "build", [1])

    def test_unfinished_external_tool_cannot_be_reported_complete(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            self.fixture(root)
            workflow.begin(root / "run", "imagegen", [1])
            result = workflow.assemble_report(root / "run", root / "evidence.json")
            self.assertEqual(result["status"], "FAIL")
            self.assertIn("execution events remain unfinished", workflow.read(root / "run/evidence-validation.json")["gateFailures"])

    def test_report_assembles_required_fields_from_actual_audits(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            self.fixture(root)
            result = workflow.assemble_report(root / "run", root / "evidence.json")
            self.assertEqual(result["status"], "PASS", workflow.read(root / "run/evidence-validation.json"))
            report = workflow.read(root / "run/qa-report.json")
            self.assertEqual(report["slideCount"], 1)
            self.assertEqual(report["visionAuditReport"], report["visualFidelityReport"])
            self.assertEqual(workflow.read(root / "run/user-copy-audit.json")["status"], "PASS")

    def test_pending_visual_observation_cannot_be_automatically_passed(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            _, review_path = self.fixture(root)
            review = workflow.read(review_path)
            review["visionAuditStatus"] = "PENDING"
            workflow.write(review_path, review)
            self.assertEqual(workflow.assemble_report(root / "run", root / "evidence.json")["status"], "FAIL")

    def test_changed_deck_requires_new_render_and_review(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            self.fixture(root)
            deck = Presentation(root / "deck.pptx")
            deck.slides[0].shapes[0].text = "被改掉的正文"
            deck.save(root / "deck.pptx")
            with self.assertRaisesRegex(ValueError, "changed after visual review"):
                workflow.assemble_report(root / "run", root / "evidence.json")

    def test_copy_audit_catches_actual_text_even_if_source_image_is_unavailable(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            task = self.create_task(root)
            deck = make_basic_pptx(root / "deck.pptx")
            task["userCopy"][0]["text"] = "用户指定的完全不同正文"
            self.assertEqual(workflow.audit_copy(task, deck, [])["status"], "FAIL")

    def test_manual_cutout_is_reported_but_not_a_level3_pass(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            self.fixture(root, {"routingIntent": {"operation": "rebuild", "deliverable": "layered"},
                               "editableBoundary": {"mustHaveAlpha": ["person"], "mustRemainIndependentImages": ["person"]}})
            Image.new("RGB", (10, 10), "white").save(root / "person.png")
            audit = audit_assets({"assets": [{"id": "person", "path": "person.png", "requiresAlpha": True}]}, root)
            workflow.write(root / "run/asset-audit.json", audit)
            result = workflow.assemble_report(root / "run", root / "evidence.json")
            self.assertEqual(result["status"], "FAIL")
            self.assertEqual(result["dimensions"]["editability"], "MANUAL_CUTOUT_REQUIRED")
            self.assertIn("required alpha is absent", " ".join(workflow.read(root / "run/evidence-validation.json")["gateFailures"]))

    def test_incremental_initialization_backs_up_without_touching_source(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = make_basic_pptx(root / "source.pptx")
            original = source.read_bytes()
            self.create_task(root, {"sourcePptx": "source.pptx", "routingIntent": {"operation": "incremental", "deliverable": "editable"}})
            self.assertEqual((root / "run/source-backup.pptx").read_bytes(), original)
            self.assertEqual(source.read_bytes(), original)
            self.assertEqual(workflow.read(root / "run/route.json")["taskMode"], "Mode E")

    def test_geometry_exceptions_require_every_real_intersection(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            deck = make_basic_pptx(root / "deck.pptx")
            workflow.write(root / "audit.json", {"outputPptxSha256": workflow.sha(deck),
                           "pages": [{"page": 1, "textFrameIntersections": [{"a": "a", "b": "b"}, {"a": "c", "b": "d"}]}]})
            item = {"page": 1, "intersectionIndex": 0, "reviewer": "fixture", "reason": "frame padding",
                    "renderObservation": "simulated exception", "status": "NO_VISIBLE_TEXT_OVERLAP"}
            report = {"outputPptx": str(deck), "textFrameIntersections": 2,
                      "auditArtifacts": {"textFrameAudit": "audit.json"}, "geometryExceptions": [item]}
            errors, gates = [], []
            validate_geometry_exceptions(report, root, errors, gates)
            self.assertTrue(gates)
            report["geometryExceptions"].append({**item, "intersectionIndex": 1})
            errors, gates = [], []
            validate_geometry_exceptions(report, root, errors, gates)
            self.assertEqual((errors, gates), ([], []))
            report["textFrameIntersections"] = 0
            errors, gates = [], []
            validate_geometry_exceptions(report, root, errors, gates)
            self.assertTrue(errors)
